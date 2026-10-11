#!/usr/bin/env python3
"""Extend BOTH full v11 illustrated masters. Never substitute an addendum.
Requires python-docx and matplotlib. PDF/PNG QA uses the documented renderer.
From the delivered package:
 python Unreal_Source/Scripts/publishing/build_cinderhold_v12.py --previous-books publishing/baselines --out-dir books
"""
from pathlib import Path
import argparse,hashlib,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle,Circle
ROOT=Path(__file__).resolve().parents[2]

def atlas(path):
    fig=plt.figure(figsize=(14,10.5),facecolor='#101c27')
    fig.text(.055,.945,'THE UNMADE   CINDERHOLD ADVENTURE',color='white',fontsize=24,weight='bold')
    fig.text(.055,.902,'Authored physical route and encounter map   |   Source schematic, not captured Unreal gameplay',color='#b6c7d4',fontsize=11)
    region=fig.add_axes([.055,.25,.255,.6],facecolor='#162532')
    region.set_xlim(-4000,2000);region.set_ylim(49800,53600);region.set_aspect('equal')
    region.add_patch(Rectangle((-3700,51100),2500,2100,fill=False,edgecolor='#dd9b59',lw=1.5))
    region.text(-3500,53400,'BLACKSLAG',color='#f5cf9e',fontsize=12,weight='bold')
    region.scatter([-2100,1370],[51130,50520],s=100,c=['#87d8cb','#f4d09b'])
    region.annotate('Varda\nbrief and report',(-2100,51130),xytext=(-3450,51070),color='white',fontsize=10)
    region.annotate('Bes household\ninspect and return',(1370,50520),xytext=(-1500,50200),color='white',fontsize=10)
    region.annotate('',xy=(1150,50520),xytext=(-1800,51040),arrowprops=dict(arrowstyle='<->',color='#a9c5da',lw=1.5))
    fig.text(.055,.365,'Reports name their source.\nResidents learn through\nactual conversations.',color='#a9c5da',fontsize=11)
    region.set_xticks([]);region.set_yticks([])
    for s in region.spines.values():s.set_color('#53697b')
    ax=fig.add_axes([.36,.2,.60,.65],facecolor='#162532')
    ax.set_xlim(-18,11);ax.set_ylim(-11,16.5);ax.set_aspect('equal')
    ax.set_xlabel('Meters east of field center',color='#b6c7d4');ax.set_ylabel('Meters north of field center',color='#b6c7d4')
    ax.tick_params(colors='#b6c7d4',labelsize=9);ax.grid(color='#304351',lw=.5,alpha=.4)
    def rect(x,y,w,h,color,alpha=1):ax.add_patch(Rectangle((x,y),w,h,color=color,alpha=alpha))
    rect(-14.5,-9,4,20,'#36746d',.42) # covered west approach
    rect(-10.825,-4.8,.25,12,'#8195a5')
    rect(-13,4.8,6,6,'#477471',.5)
    rect(-9,3.45,9,2,'#6f6c50',.5)
    rect(-16.5,8,7,2,'#5b657e',.5)
    # Gate and enclosed archive geometry matches source coordinates.
    rect(-16.7,7.625,3.4,.35,'#8195a5');rect(-16.7,10.025,3.4,.35,'#8195a5')
    rect(-16.875,7.65,.35,2.7,'#8195a5')
    rect(-7.2,3.125,3.6,.35,'#8195a5');rect(-7.2,5.525,3.6,.35,'#8195a5')
    rect(-7.175,3.25,.35,2.4,'#efc36b');rect(-13.675,7.8,.35,2.4,'#efc36b')
    ax.add_patch(Circle((0,6),9.5,fill=False,ls='--',lw=1.2,edgecolor='#b97168'))
    ax.add_patch(Circle((0,6),4.2,fill=False,ls=':',lw=1.2,edgecolor='#ffac78'))
    ax.scatter([0],[6],s=150,marker='D',color='#f09b70')
    ax.text(-3,13.8,'Sootbound Elder',color='white',fontsize=11,weight='bold')
    ax.text(-3,12.1,'Territory 9.5 m\nAttack intention within 4.2 m',color='#dfbda9',fontsize=9)
    for (x,y),label in zip([(-12.5,-4.5),(-12.5,0),(-12.5,4.5)],['Tracks','Service log','Cadence plate']):
        ax.scatter([x],[y],s=55,color='#b1e4ce');ax.annotate(label,(x,y),xytext=(-16.7,y-.4),color='white',fontsize=9)
    ax.scatter([-8,-8,0],[6,9,-8.7],s=[80,80,65],color=['#efc36b','#cdb8ef','#87d8cb'])
    ax.text(-9.8,5.2,'Vent',color='white',fontsize=10);ax.text(-9.9,10,'Cadence post',color='white',fontsize=10)
    ax.text(.6,-9.3,'Varda',color='white',fontsize=10)
    ax.text(-16.1,11,'Archive alcove',color='#c0cce0',fontsize=9)
    ax.text(-6.8,2.5,'Kiln passage',color='#d9cfac',fontsize=9)
    for s in ax.spines.values():s.set_color('#53697b')
    fig.text(.055,.12,'E read and maintain   F Glimpse   Q Fold   R/T preview permanent route choices',color='white',fontsize=12)
    fig.text(.055,.075,'Two elder outcomes × three routes. Varda report → Bes report → saved return at the household shelf.',color='#b6c7d4',fontsize=11)
    fig.text(.055,.04,'Wages, warmth and consent keep their existing authority. Engine navigation, LOS, staging and import remain to verify.',color='#b6c7d4',fontsize=10)
    fig.savefig(path,dpi=160,facecolor=fig.get_facecolor());plt.close(fig)

def table(doc,lines):
    records=[[cell.strip() for cell in line.strip('|').split('|')] for line in lines if not line.startswith('| ---')]
    t=doc.add_table(rows=1,cols=len(records[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.style='Table Grid';t.autofit=False
    widths=[1.35,1.45,1.0,2.4]
    for i,width in enumerate(widths):t.columns[i].width=Inches(width)
    normal_size=doc.styles["Normal"].font.size.pt
    table_size=min(10.0,normal_size)
    for i,title in enumerate(records[0]):t.rows[0].cells[i].text=title
    for record in records[1:]:
        cells=t.add_row().cells
        for i,value in enumerate(record):cells[i].text=value
    for ri,row in enumerate(t.rows):
        for i,cell in enumerate(row.cells):
            cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for side in ('top','bottom','left','right'):
                v=OxmlElement('w:'+side);v.set(qn('w:w'),'85');v.set(qn('w:type'),'dxa');m.append(v)
            pr.append(m)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E5EAF0');pr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for side in ('top','bottom','left','right'):
                b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            pr.append(borders)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(3)
                for run in p.runs:run.font.size=Pt(table_size);run.font.bold=ri==0;run.font.color.rgb=RGBColor(0,0,0)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph()

def main():
    p=argparse.ArgumentParser();p.add_argument('--source-root',type=Path,default=ROOT)
    p.add_argument('--previous-books',type=Path,required=True);p.add_argument('--out-dir',type=Path,required=True);a=p.parse_args()
    a.out_dir.mkdir(parents=True,exist_ok=True)
    art=a.out_dir/'THE_UNMADE_v12_Cinderhold_Adventure_Atlas.png';atlas(art)
    chapter=(a.source_root/'docs/lore/CHAPTER_32_CINDERHOLD_ADVENTURE.md').read_text()
    inputs=[]
    for prefix in ['THE_UNMADE_Complete_Illustrated_Lore_Bible_Volume_X','THE_UNMADE_Volume_X_Illustrated']:
        previous=a.previous_books/f'{prefix}_Paths_and_Predators_v11.docx'
        doc=Document(previous);old=len(doc.inline_shapes);doc.add_page_break()
        lines=chapter.splitlines();index=0
        while index<len(lines):
            line=lines[index];index+=1
            if not line.strip():continue
            if line.startswith('|'):
                data=[line]
                while index<len(lines) and lines[index].startswith('|'):data.append(lines[index]);index+=1
                table(doc,data);continue
            if line.startswith('## '):
                heading=doc.add_heading(line[3:],1)
                for run in heading.runs:run.font.color.rgb=RGBColor(0,0,0)
            elif line.startswith('### '):
                heading=doc.add_heading(line[4:],2)
                for run in heading.runs:run.font.color.rgb=RGBColor(0,0,0)
                if line.startswith('### 32.1'):
                    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    para.add_run().add_picture(str(art),width=Inches(6.3))
                    doc.add_paragraph('Figure X 32 Cinderhold adventure and Blackslag route map. Exact source positions; Unreal playtesting outstanding.',style='Caption')
            else:
                para=doc.add_paragraph(line)
                para.paragraph_format.line_spacing=1.14
                para.paragraph_format.space_after=Pt(4.5)
        doc.core_properties.title='The Unmade Cinderhold Adventure v12'
        target=a.out_dir/f'{prefix}_Cinderhold_Adventure_v12.docx';doc.save(target)
        assert len(doc.inline_shapes)==old+1
        inputs.append({'name':previous.name,'sha256':hashlib.sha256(previous.read_bytes()).hexdigest(),'retained_images':old})
        print(target.name,len(doc.inline_shapes),'images')
    (a.out_dir/'v12_build_inputs.json').write_text(json.dumps(inputs,indent=2)+'\n')
if __name__=='__main__':main()
