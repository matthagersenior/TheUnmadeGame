"""Extend both intact v12 masters, preserving every inherited paragraph/image."""
from pathlib import Path
import argparse,hashlib,json,textwrap
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
ROOT=Path(__file__).resolve().parents[2]
def atlas(path):
 d=json.loads((ROOT/'Authoring/remembered_roads_v13.json').read_text());fig=plt.figure(figsize=(8.2,6.7),facecolor='#11212a')
 fig.text(.055,.95,'THE UNMADE   REMEMBERED ROADS',fontsize=17,weight='bold',color='white')
 fig.text(.055,.91,'Nine realms  /  Eleven shipments  /  Eighty-eight personal returns',fontsize=10.5,color='#c2d4d9')
 for i,c in enumerate(d['cities']):
  x=.05+(i%2)*.46;y=.775-(i//2)*.125
  ax=fig.add_axes([x,y,.43,.115],facecolor='#1b343e');ax.set_xticks([]);ax.set_yticks([])
  for sp in ax.spines.values():sp.set_color('#57747a')
  ax.text(.045,.79,c['id'],fontsize=12,color='#efc993',weight='bold',transform=ax.transAxes)
  ax.text(.045,.48,'\n'.join(textwrap.wrap(c['cargo'],37)),fontsize=10.5,color='white',transform=ax.transAxes,va='center')
  ax.text(.045,.12,'Receive at '+d['cities'][c['destination']]['id'],fontsize=10.5,color='#b4ddd4',transform=ax.transAxes)
 fig.text(.53,.17,'44 finite sealed units\nNo off-screen deliveries\nExisting travel remains authoritative',fontsize=11,color='#dce6e8',linespacing=1.55)
 fig.text(.055,.035,'Authored source ledger  /  Unreal compilation and playtesting deferred',fontsize=9.5,color='#c2d4d9')
 fig.savefig(path,dpi=200,facecolor=fig.get_facecolor());plt.close(fig)
def add_table(doc,lines):
 rows=[[x.strip() for x in l.strip('|').split('|')] for l in lines if not l.startswith('| ---')];t=doc.add_table(rows=0,cols=4);t.style='Table Grid';t.autofit=False;w=[1.,2.25,1.75,1.2]
 t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for j,width in enumerate(w):t.columns[j].width=Inches(width)
 for ri,record in enumerate(rows):
  row=t.add_row()
  for j,value in enumerate(record):
   cell=row.cells[j];cell.width=Inches(w[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cell.text=value;pr=cell._tc.get_or_add_tcPr()
   margins=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    m=OxmlElement('w:'+side);m.set(qn('w:w'),'85');m.set(qn('w:type'),'dxa');margins.append(m)
   pr.append(margins);b=OxmlElement('w:tcBorders')
   for side in ['top','bottom','left','right']:
    x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
   pr.append(b)
   if ri==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DFE8EC');pr.append(shade)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(3);para.paragraph_format.line_spacing=1.1
    for run in para.runs:run.font.size=Pt(min(10,doc.styles['Normal'].font.size.pt));run.bold=ri==0
  if ri==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
def main():
 p=argparse.ArgumentParser();p.add_argument('--previous-books',type=Path,required=True);p.add_argument('--out-dir',type=Path,required=True);a=p.parse_args();a.out_dir.mkdir(parents=True,exist_ok=True)
 art=a.out_dir/'THE_UNMADE_v13_Remembered_Roads_Atlas.png';atlas(art);source=(ROOT/'docs/lore/CHAPTER_33_REMEMBERED_ROADS.md').read_text().splitlines();inputs=[]
 for prefix in ['THE_UNMADE_Complete_Illustrated_Lore_Bible_Volume_X','THE_UNMADE_Volume_X_Illustrated']:
  previous=a.previous_books/(prefix+'_Cinderhold_Adventure_v12.docx');doc=Document(previous);old=len(doc.inline_shapes);doc.add_page_break();i=0
  while i<len(source):
   line=source[i];i+=1
   if not line.strip():continue
   if line.startswith('|'):
    lines=[line]
    while i<len(source) and source[i].startswith('|'):lines.append(source[i]);i+=1
    add_table(doc,lines);continue
   if line.startswith('## '):h=doc.add_heading(line[3:],1)
   elif line.startswith('### '):
    h=doc.add_heading(line[4:],2)
    if line.startswith('### 33.1'):
     para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.add_run().add_picture(str(art),width=Inches(6.3));doc.add_paragraph('Figure X 33 The eleven physical shipments and receiving cities. Source ledger; engine testing deferred.',style='Caption')
   else:
    para=doc.add_paragraph(line);para.paragraph_format.line_spacing=1.14;para.paragraph_format.space_after=Pt(4.5);continue
   for run in h.runs:run.font.color.rgb=RGBColor(0,0,0)
  doc.core_properties.title='The Unmade World of Remembered Roads v13';target=a.out_dir/(prefix+'_Remembered_Roads_v13.docx');doc.save(target);assert len(doc.inline_shapes)==old+1
  inputs.append({'name':previous.name,'sha256':hashlib.sha256(previous.read_bytes()).hexdigest(),'retained_images':old});print(target.name,len(doc.inline_shapes))
 (a.out_dir/'v13_build_inputs.json').write_text(json.dumps(inputs,indent=2)+'\n')
if __name__=='__main__':main()
